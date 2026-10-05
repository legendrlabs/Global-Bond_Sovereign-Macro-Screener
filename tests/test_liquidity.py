from datetime import date
from io import BytesIO
import json
import unittest

from openpyxl import Workbook
from sovereign_macro import market_inputs
from sovereign_macro.config import load_config
from sovereign_macro.http import Payload
from sovereign_macro.models import DataError
from sovereign_macro.pipeline import demo_bundle, evaluate
from sovereign_macro.summary import build_executive_summary, render_summary_text


def workbook(points=None, unit='(100 million yen)'):
    w=Workbook();s=w.active;s.title='(Ｂ)一般売買高'
    s.cell(1,1,unit);s.cell(1,4,'Trading volume（Sell＋Purchase）')
    s.cell(2,4,'一般売買\nOutright transactions');s.cell(3,4,'国債\nGovernment Bonds')
    for i,(period,value) in enumerate(points or [('2026/07',120),('2026/08',150)],5):
        s.cell(i,1,period);s.cell(i,2,'合計');s.cell(i,3,'Total');s.cell(i,4,value)
    b=BytesIO();w.save(b);return b.getvalue()


def survey(points):
    return json.dumps({'pd':{'timeseries':points}}).encode()


class LiquidityTests(unittest.TestCase):
    def parser(self,name):
        fn=getattr(market_inputs,name,None)
        self.assertTrue(callable(fn),'Missing liquidity parser: '+name)
        return fn

    def test_nyfed_preserves_selected_series_date_and_daily_average_unit(self):
        parse=self.parser('parse_nyfed_liquidity')
        body=survey([{'keyid':'PDGSWOEXTTOT','asofdate':'2026-09-23','value':'100'},
                     {'keyid':'OTHER','asofdate':'2026-10-01','value':'900'}])
        row=parse(body,date(2026,10,6))
        self.assertEqual(row['value'],100)
        self.assertEqual(row['observation_date'],'2026-09-23')
        self.assertEqual(row['unit'],'USD_million_daily_average')
        self.assertFalse(row['usable_for_scoring'])

    def test_nyfed_missing_invalid_duplicates_and_future_are_not_current_values(self):
        parse=self.parser('parse_nyfed_liquidity')
        point={'keyid':'PDGSWOEXTTOT','asofdate':'2026-09-23','value':'100'}
        for val in ('*','NaN','-1'):
            with self.subTest(value=val),self.assertRaises(DataError):
                parse(survey([dict(point,value=val)]),date(2026,10,6))
        with self.assertRaises(DataError): parse(survey([point,point]),date(2026,10,6))
        row=parse(survey([point,dict(point,asofdate='2026-11-01',value='200')]),date(2026,10,6))
        self.assertEqual(row['value'],100)
        row=parse(survey([point]),date(2026,12,6))
        self.assertEqual(row['status'],'STALE');self.assertIsNone(row['value'])

    def test_jsda_uses_reported_outright_total_without_daily_conversion(self):
        row=self.parser('parse_jsda_liquidity')(workbook(),date(2026,10,6))
        self.assertEqual(row['value'],150)
        self.assertEqual(row['period'],'2026-08')
        self.assertEqual(row['observation_date'],'2026-08-31')
        self.assertEqual(row['unit'],'JPY_100_million_monthly_sell_plus_purchase')
        self.assertFalse(row['usable_for_scoring'])

    def test_jsda_rejects_unit_drift_conflicting_totals_and_nonfinite_value(self):
        parse=self.parser('parse_jsda_liquidity')
        for body in [workbook(unit='JPY billion'),workbook([('2026/08',100),('2026/08',200)]),
                     workbook([('2026/08','NaN')]),b'not an Excel workbook']:
            with self.subTest(length=len(body)),self.assertRaises(DataError): parse(body,date(2026,10,6))

    def test_jsda_completed_month_rule_and_structural_freshness(self):
        parse=self.parser('parse_jsda_liquidity')
        row=parse(workbook([('2026/08',150),('2026/10',200)]),date(2026,10,6))
        self.assertEqual(row['value'],150)
        row=parse(workbook(),date(2027,3,1))
        self.assertEqual(row['status'],'STALE');self.assertIsNone(row['value'])

    def test_optional_sources_fail_independently_and_keep_provenance(self):
        class Client:
            def fetch(self,url):
                if 'stats.bis.org' in url: raise TimeoutError('BIS offline')
                if '/list/timeseries' in url:
                    body=json.dumps({'pd':{'timeseries':[{'seriesbreak':'SBN2024','keyid':'PDGSWOEXTTOT',
                        'description':'Total - U.S. TREASURY SECURITIES (EXCLUDING TREASURY INFLATION-PROTECTED SECURITIES (TIPS)) - DEALER TRANSACTIONS WITH INTER-DEALER BROKERS + Total - U.S. TREASURY SECURITIES (EXCLUDING TREASURY INFLATION-PROTECTED SECURITIES (TIPS)) - DEALER TRANSACTIONS WITH OTHERS'}]}}).encode()
                elif '/api/pd/' in url: raise TimeoutError('USA offline')
                elif url.endswith('koushasai.xlsx'): body=workbook()
                else: raise AssertionError('unexpected source '+url)
                return Payload(body,url,'2026-10-05T15:30:00Z','hash','', 'curl_cffi')
        result=market_inputs.collect_market_inputs(Client(),{'enabled':True,'liquidity_enabled':True},date(2026,10,6))
        self.assertEqual(result.get('liquidity',{}).get('JPN',{}).get('value'),150)
        self.assertEqual(result['liquidity']['USA']['status'],'UNAVAILABLE')
        self.assertEqual(result['liquidity']['JPN']['raw_sha256'],'hash')
        self.assertEqual(result['liquidity']['JPN']['transport'],'curl_cffi')

    def test_liquidity_is_displayed_but_never_creates_score_and_public_or_demo_hide_value(self):
        config=load_config();day=date(2026,10,6);bundle=demo_bundle(config,day)
        bundle['market_inputs']={'size':{},'liquidity':{'JPN':dict(value=876.543,provider='jsda',
            status='AVAILABLE_DIAGNOSTIC',period='2026-08',unit='JPY_100_million_monthly_sell_plus_purchase',
            redistribution='pending',raw_sha256='hash',definition='member outright sell plus purchase')},'errors':[]}
        private=evaluate(config,bundle,day)
        row=next(r for r in private['rows'] if r['iso3']=='JPN')
        self.assertEqual(row['market_inputs']['liquidity']['value'],876.543)
        self.assertIsNone(row['market_quality_norm']);self.assertIsNone(row['adjusted'])
        self.assertIn('876.543',render_summary_text(build_executive_summary(private)))
        for options in ({'public_output':True},{'demo':True}):
            result=evaluate(config,bundle,day,**options)
            self.assertNotIn('876.543',json.dumps(result))
        self.assertEqual(bundle['market_inputs']['liquidity']['JPN']['value'],876.543)

    def test_whole_public_row_gate_hides_otherwise_cleared_liquidity(self):
        config=load_config();day=date(2026,10,6);bundle=demo_bundle(config,day)
        bundle['market_inputs']={'liquidity':{'JPN':dict(value=543.21,status='AVAILABLE_DIAGNOSTIC',
            redistribution='allowed',raw_sha256='hash')},'errors':[]}
        result=evaluate(config,bundle,day,public_output=True)
        self.assertNotIn('543.21',json.dumps(result))

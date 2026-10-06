from datetime import date, datetime
from io import BytesIO
import unittest
from openpyxl import Workbook
from sovereign_macro.europe_liquidity import (parse_dmo_liquidity, parse_riksbank_liquidity,
    verify_riksbank_description, discover_workbook, collect_europe_liquidity,
    DMO_PAGE, RIKSBANK_PAGE, RIKSBANK_DESCRIPTION)
from sovereign_macro.http import Payload
from sovereign_macro.models import DataError
from sovereign_macro.market_inputs import country_market_inputs
from sovereign_macro.config import load_config
from sovereign_macro.pipeline import demo_bundle, evaluate


def save(book):
    out=BytesIO();book.save(out);return out.getvalue()


def dmo(change=None):
    w=Workbook();s=w.active;s.title='2026'
    for row in [('ALL GEMMs turnover by maturity band',),(2026,),('Q2 (April - June)',),
                ('BAND','Customer (£ million)','Customer (% Aggregate)','Professional (£ million)',
                 'Professional (% Aggregate)','Combined (£ million)','Combined (% Aggregate)'),
                ('3-5 years',10,0,20,0,30,0),('Index-linked',5,0,10,0,15,0),('Total',15,0,30,0,45,0)]: s.append(row)
    n=w.create_sheet('Notes')
    n.append(('market values; sum of all trade reports; neither counterparty is a GEMM; '
              'total purchases and sales; do not include turnover on repo',))
    if change: change(w)
    return save(w)


def sweden(rows=None,header=None):
    w=Workbook();s=w.active;s.title='Data'
    s.append(header or ('Trade date','Asset','Contract','Counterparty','Amount (million SEK)',
                      'Adjusted amount (million SEK)','Complete'))
    for row in rows or [(datetime(2026,9,25),'GVB','SP','REP',20,10,True),
                        (datetime(2026,9,25),'GVB','SP','CUSE',30,30,True),
                        (datetime(2026,9,25),'GVB','SP','PRIMM',900,900,True),
                        (datetime(2026,9,25),'GVB','RE','CUSE',800,800,True),
                        (datetime(2026,9,25),'ILB','SP','CUSE',700,700,True)]: s.append(row)
    return save(w)


def description(change=None):
    w=Workbook();s=w.active;s.title='Description'
    for k,v in [('Unit','Millions SEK'),('GVB','Government Bonds'),('SP','Spot'),('RE','Repo'),
                ('PRIMM','Secondary Market is the sum of all other counterparts'),
                ('Adjusted amount','Adjusted for double-counting; dividing those amounts by two')]: s.append((None,k,v))
    if change: change(w)
    return save(w)


class EuropeLiquidityTests(unittest.TestCase):
    def test_dmo_reports_quarter_total_without_halving_or_daily_conversion(self):
        row=parse_dmo_liquidity(dmo(),date(2026,10,6))
        self.assertEqual(row['value'],45);self.assertEqual(row['period'],'2026-Q2')
        self.assertEqual(row['observation_date'],'2026-06-30');self.assertFalse(row['usable_for_scoring'])
        self.assertIn('index-linked',row['definition'])

    def test_dmo_rejects_total_definition_unit_and_duplicate_drift(self):
        changes=[lambda w:setattr(w['2026']['F7'],'value',44),
                 lambda w:setattr(w['2026']['B4'],'value','Customer (EUR million)'),
                 lambda w:w.remove(w['Notes']),
                 lambda w:w['2026'].append(('Total',15,0,30,0,45,0))]
        for change in changes:
            with self.subTest(change=change),self.assertRaises(DataError):
                parse_dmo_liquidity(dmo(change),date(2026,10,6))

    def test_dmo_ignores_unrelated_old_sheet_and_has_separate_age_policy(self):
        row=parse_dmo_liquidity(dmo(lambda w:w.create_sheet('2001')),date(2026,10,6))
        self.assertEqual(row['value'],45)
        self.assertEqual(parse_dmo_liquidity(dmo(),date(2027,2,1))['status'],'STALE')
        with self.assertRaises(DataError):parse_dmo_liquidity(dmo(),date(2026,6,1))

    def test_sweden_uses_adjusted_secondary_spot_only_and_separate_freshness(self):
        row=parse_riksbank_liquidity(sweden(),date(2026,9,30))
        self.assertEqual(row['value'],40);self.assertFalse(row['usable_for_scoring'])
        self.assertEqual(row['unit'],'SEK_million_daily_adjusted_turnover')
        stale=parse_riksbank_liquidity(sweden(),date(2026,10,6))
        self.assertEqual(stale['status'],'STALE');self.assertIsNone(stale['value'])

    def test_sweden_latest_incomplete_cannot_fall_back_to_older_complete(self):
        rows=[(datetime(2026,9,24),'GVB','SP','REP',20,10,True),
              (datetime(2026,9,25),'GVB','SP','REP',20,10,False)]
        with self.assertRaisesRegex(DataError,'INCOMPLETE_LATEST'):
            parse_riksbank_liquidity(sweden(rows),date(2026,10,6))

    def test_sweden_future_duplicate_bad_unit_and_invalid_values(self):
        point=(datetime(2026,9,25),'GVB','SP','REP',20,10,True)
        future=(datetime(2026,11,1),'GVB','SP','REP',40,20,True)
        self.assertEqual(parse_riksbank_liquidity(sweden([point,future]),date(2026,9,30))['value'],10)
        for rows in [[point,point],[point[:-2]+('NaN',True)],
                     [point[:5]+(25,True)],[point[:-1]+(None,)]]:
            with self.subTest(rows=rows),self.assertRaises(DataError):
                parse_riksbank_liquidity(sweden(rows),date(2026,9,30))
        with self.assertRaises(DataError):parse_riksbank_liquidity(sweden(header=('bad',)),date(2026,9,30))

    def test_definition_metadata_is_verified(self):
        verify_riksbank_description(description())
        with self.assertRaises(DataError):
            verify_riksbank_description(description(lambda w:setattr(w.active['C2'],'value','Corporate Bonds')))

    def test_discovery_rejects_host_drift_multiple_files_and_wrong_format(self):
        label="All GEMMs' turnover by maturity band"
        body=f'<a href="/media/latest.xlsx">{label}</a>'.encode()
        self.assertEqual(discover_workbook(body,DMO_PAGE,label),'https://www.dmo.gov.uk/media/latest.xlsx')
        for body in [f'<a href="https://example.com/a.xlsx">{label}</a>',
                     f'<a href="/a.xlsx">{label}</a><a href="/b.xlsx">{label}</a>',
                     f'<a href="/a.xls">{label}</a>']:
            with self.assertRaises(DataError):discover_workbook(body.encode(),DMO_PAGE,label)

    def test_collect_isolates_failures_preserves_provenance_and_respects_flags(self):
        class Client:
            def __init__(self):self.calls=[]
            def fetch(self,url):
                self.calls.append(url)
                if url==DMO_PAGE:raise TimeoutError('UK offline')
                if url==RIKSBANK_PAGE:body=b'<a href="/contentassets/latest.xlsx">Fixed Income, Daily Turnover</a>'
                elif url==RIKSBANK_DESCRIPTION:body=description()
                else:body=sweden()
                return Payload(body,url,'2026-09-30T10:00:00Z','hash')
        client=Client();rows=collect_europe_liquidity(client,{},date(2026,9,30))
        self.assertEqual(rows['GBR']['status'],'UNAVAILABLE');self.assertEqual(rows['SWE']['value'],40)
        self.assertEqual(rows['SWE']['definition_sha256'],'hash')
        self.assertEqual(rows['SWE']['index_url'],RIKSBANK_PAGE)
        client=Client();self.assertEqual(collect_europe_liquidity(client,dict(uk_dmo_enabled=False,riksbank_enabled=False),date(2026,9,30)),{})
        self.assertFalse(client.calls)

    def test_country_output_public_and_demo_do_not_leak_amounts(self):
        bundle={'market_inputs':{'liquidity':{'GBR':dict(value=123.4567,provider='uk_dmo',
            status='AVAILABLE_DIAGNOSTIC',redistribution='pending')}}}
        self.assertEqual(country_market_inputs('GBR',bundle)['liquidity']['value'],123.4567)
        for opts in [dict(demo=True),dict(public_output=True)]:
            self.assertNotIn('123.4567',str(country_market_inputs('GBR',bundle,**opts)))

    def test_both_providers_leave_adjusted_unavailable_and_mask_provenance(self):
        config=load_config();day=date(2026,9,30);bundle=demo_bundle(config,day)
        uk=parse_dmo_liquidity(dmo(),day);swe=parse_riksbank_liquidity(sweden(),day)
        for item in (uk,swe):item.update(raw_sha256='source-hash',redistribution='pending')
        bundle['market_inputs']={'liquidity':{'GBR':uk,'SWE':swe}}
        private=evaluate(config,bundle,day)
        for row in private['rows']:
            if row['iso3'] not in ('GBR','SWE'):continue
            self.assertIsNotNone(row['market_inputs']['liquidity']['value'])
            self.assertIsNone(row['adjusted']);self.assertIsNone(row['market_quality_norm'])
        for opts in ({'public_output':True},{'demo':True}):
            result=evaluate(config,bundle,day,**opts)
            for row in result['rows']:
                if row['iso3'] not in ('GBR','SWE'):continue
                self.assertIsNone(row['market_inputs']['liquidity'].get('value'))
                for item in row['provenance']:
                    if item.get('provider') in ('uk_dmo','riksbank'):
                        self.assertIsNone(item.get('value'))


if __name__=='__main__':unittest.main()

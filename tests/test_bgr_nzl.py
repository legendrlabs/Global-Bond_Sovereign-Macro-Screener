from copy import deepcopy
from datetime import date, datetime, timedelta
from io import BytesIO
import unittest
from zipfile import ZipFile
import re
from openpyxl import Workbook
from sovereign_macro.rbnz import parse_rbnz, RBNZ_WORKBOOK, RBNZ_SERIES
from sovereign_macro.fx import redenominated_points, fx_metrics
from sovereign_macro.models import DataError
from sovereign_macro.http import Payload
from sovereign_macro.config import load_config
from sovereign_macro.pipeline import demo_bundle, evaluate
from sovereign_macro.yield_selection import collect_preferred_yield
from sovereign_macro.sources import collect_yield
from sovereign_macro.summary import build_executive_summary, render_summary_markdown, render_summary_html


TRANSITION=dict(from_currency='BGN',to_currency='EUR',effective_date='2026-01-01',
                old_units_per_new_unit=1.95583,evidence_url='https://www.ecb.europa.eu/euro/changeover/bulgaria/html/index.en.html')


def xlsx(change=None, bad_dimensions=False):
    w=Workbook();s=w.active;s.title='Data'
    group='Secondary market government bond closing yields'
    for row in [(None,group),(None,'5 year'),('Notes',None),('Unit','%pa'),('Series Id',RBNZ_SERIES),
                (datetime(2026,10,2),4.4),(datetime(2026,10,5),4.5)]:s.append(row)
    s=w.create_sheet('Table Description')
    for row in [('Published By','Reserve Bank of New Zealand'),('Table','Daily wholesale interest rates (% pa) - B2'),
                ('Published Date',datetime(2026,10,6))]:s.append(row)
    s=w.create_sheet('Series Definitions');s.append(('Group','Series','Series Id','Unit','Note'))
    s.append((group,'5 year',RBNZ_SERIES,'%pa',None))
    if change:change(w)
    b=BytesIO();w.save(b)
    if not bad_dimensions:return b.getvalue()
    out=BytesIO()
    with ZipFile(BytesIO(b.getvalue())) as src,ZipFile(out,'w') as dst:
        for name in src.namelist():
            body=src.read(name)
            if name.startswith('xl/worksheets/'):
                body=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',body)
            dst.writestr(name,body)
    return out.getvalue()


class BgrNzlTests(unittest.TestCase):
    def test_rbnz_recovers_supplier_dimension_bug_and_keeps_dates(self):
        rows,pub=parse_rbnz(xlsx(bad_dimensions=True))
        self.assertEqual(rows[-1],('2026-10-05',4.5));self.assertEqual(pub,'2026-10-06')

    def test_rbnz_rejects_unit_tenor_duplicate_and_invalid_value(self):
        for change in [lambda w:setattr(w['Data']['B4'],'value','bps'),
                       lambda w:setattr(w['Data']['B2'],'value','10 year'),
                       lambda w:w['Data'].append((datetime(2026,10,5),4.6)),
                       lambda w:setattr(w['Data']['B7'],'value','NaN'),
                       lambda w:setattr(w['Series Definitions']['B2'],'value','Swap')]:
            with self.subTest(change=change),self.assertRaises(DataError):parse_rbnz(xlsx(change))

    def test_rbnz_metadata_and_missing_values_not_invented(self):
        rows,_=parse_rbnz(xlsx(lambda w:setattr(w['Data']['B7'],'value','..')))
        self.assertEqual(rows[-1][0],'2026-10-02')
        with self.assertRaises(DataError):parse_rbnz(xlsx(),series='other')
        with self.assertRaises(DataError):parse_rbnz(xlsx(lambda w:setattr(w['Table Description']['B3'],'value','today')))
        with self.assertRaises(DataError):parse_rbnz(xlsx(lambda w:setattr(w['Data']['A7'],'value',datetime(2026,10,7))))

    def test_official_future_publication_stale_date_and_redirect_fail(self):
        config=load_config();country=next(c for c in config['countries']['countries'] if c['iso3']=='NZL')
        class Client:
            def __init__(self,body,url=RBNZ_WORKBOOK):self.body=body;self.url=url
            def fetch(self,*args,**kwargs):return Payload(self.body,self.url,'2026-10-06T10:00:00Z','hash')
        with self.assertRaisesRegex(DataError,'FUTURE_PUBLICATION'):
            collect_yield(Client(xlsx()),country,date(2026,10,5))
        with self.assertRaisesRegex(DataError,'REDIRECT'):
            collect_yield(Client(xlsx(),'https://example.com/other.xlsx'),country,date(2026,10,6))
        with self.assertRaises(DataError):
            obs=collect_yield(Client(xlsx()),country,date(2026,10,20));obs.valid_on(date(2026,10,20),7)

    def test_official_selected_before_wgb_and_benchmark_warning_preserved(self):
        class Client:
            def __init__(self):self.calls=[]
            def fetch(self,url,**kwargs):
                self.calls.append(url);return Payload(xlsx(),url,'2026-10-06T10:00:00Z','hash')
        config=load_config();country=next(c for c in config['countries']['countries'] if c['iso3']=='NZL')
        client=Client();obs=collect_preferred_yield(client,country,config,date(2026,10,6))
        self.assertEqual(client.calls,[RBNZ_WORKBOOK]);self.assertEqual(obs.value,4.5)
        self.assertEqual(obs.period,'2026-10-05');self.assertEqual(obs.source_date,'2026-10-06')
        self.assertEqual(obs.yield_type,'benchmark');self.assertEqual(obs.selection_reason,'OFFICIAL_VALID')
        bundle=demo_bundle(config,date(2026,10,6));bundle['yields']['NZL:5']=obs
        row=next(r for r in evaluate(config,bundle,date(2026,10,6))['rows'] if r['iso3']=='NZL')
        self.assertTrue(row['usable_baseline']);self.assertIn('BASELINE_DEFINITION_DIFFERENCE',row['warnings'])

    def test_conversion_direction_boundary_and_rounding_are_exact(self):
        points={date(2025,12,31):{'KRW':1500,'BGN':1.9558},date(2026,1,2):{'KRW':1500}}
        result=redenominated_points(points,'EUR',date(2026,10,6),TRANSITION)
        self.assertAlmostEqual(result[date(2025,12,31)]['KRW'],1500/1.9558*1.95583)
        self.assertEqual(result[date(2026,1,2)]['KRW'],1500)
        self.assertEqual(points[date(2025,12,31)]['BGN'],1.9558)

    def test_missing_predecessor_is_not_historical_eur_and_bad_contract_fails(self):
        points={date(2025,12,31):{'KRW':1500},date(2026,1,2):{'KRW':1500}}
        self.assertEqual(redenominated_points(points,'EUR',date(2026,10,6),TRANSITION)[date(2025,12,31)],{})
        for change in [dict(old_units_per_new_unit=0),dict(old_units_per_new_unit=float('nan')),
                       dict(to_currency='USD'),dict(effective_date='bad')]:
            with self.subTest(change=change),self.assertRaises(DataError):
                redenominated_points(points,'EUR',date(2026,10,6),dict(TRANSITION,**change))
        with self.assertRaises(DataError):redenominated_points(points,'EUR',date(2025,12,31),TRANSITION)

    def test_redenominated_history_retains_gap_sample_and_stale_gates(self):
        start=date(2023,10,1);asof=date(2026,10,6)
        points={start+timedelta(days=i):dict(KRW=1500+i/100,BGN=1.9558) for i in range((asof-start).days+1)
                if (start+timedelta(days=i)).weekday()<5}
        fx=fx_metrics(points,'EUR',asof,transition=TRANSITION)
        self.assertIsNotNone(fx['vol_1y']);self.assertIsNotNone(fx['vol_3y']);self.assertFalse(fx['errors'])
        for d in list(points):
            if date(2025,5,1)<=d<=date(2025,5,20):points[d].pop('BGN')
        fx=fx_metrics(points,'EUR',asof,transition=TRANSITION)
        self.assertIn('FX_GAP',fx['errors']);self.assertIsNone(fx['vol_3y'])
        stale=fx_metrics(points,'EUR',date(2026,11,6),transition=TRANSITION)
        self.assertIn('FX_STALE',stale['errors'])

    def test_pipeline_bgr_exempts_only_verified_transition_history(self):
        config=load_config();day=date(2026,10,6);bundle=demo_bundle(config,day)
        for rates in bundle['fx'].values():rates['BGN']=1.9558
        row=next(r for r in evaluate(config,bundle,day)['rows'] if r['iso3']=='BGR')
        self.assertIsNotNone(row['fx_vol_1y']);self.assertIsNotNone(row['fx_vol_3y'])
        self.assertIn('FX_REDENOMINATED_HISTORY',row['warnings'])
        self.assertNotIn('CURRENCY_1Y_HISTORY_UNAVAILABLE',row['errors'])
        self.assertIsNone(row['adjusted'])
        broken=deepcopy(config);country=next(c for c in broken['countries']['countries'] if c['iso3']=='BGR')
        country['fx_transition']['effective_date']='2025-01-01'
        row=next(r for r in evaluate(broken,bundle,day)['rows'] if r['iso3']=='BGR')
        self.assertIsNone(row['fx_vol_1y'])

    def test_public_report_explains_transition_without_unmasking_numbers(self):
        config=load_config();day=date(2026,10,6);bundle=demo_bundle(config,day)
        for rates in bundle['fx'].values():rates['BGN']=1.9558
        result=evaluate(config,bundle,day,public_output=True)
        row=next(r for r in result['rows'] if r['iso3']=='BGR')
        self.assertIsNone(row['fx_vol_1y']);self.assertIsNone(row['fx_vol_3y'])
        summary=build_executive_summary(result)
        for rendered in (render_summary_markdown(summary),render_summary_html(summary)):
            self.assertIn('FX History Definitions',rendered)
            self.assertIn('not historical EUR observations',rendered)


if __name__=='__main__':unittest.main()

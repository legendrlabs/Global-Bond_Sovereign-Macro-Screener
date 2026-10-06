"""RBNZ B2 daily close benchmark, retaining trade and publication dates."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from .models import DataError, finite

RBNZ_WORKBOOK='https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx'
RBNZ_SERIES='INM.DG105.NZZCF'
RBNZ_DEFINITION='https://www.rbnz.govt.nz/en/statistics/series/exchange-and-interest-rates/wholesale-interest-rates'


def parse_rbnz(body, series=RBNZ_SERIES):
    if series!=RBNZ_SERIES: raise DataError('RBNZ_SERIES_CONTRACT')
    try:
        with ZipFile(BytesIO(body)) as archive:
            if len(archive.infolist())>1000 or sum(i.file_size for i in archive.infolist())>30_000_000:
                raise DataError('RBNZ_WORKBOOK_TOO_LARGE')
        book=load_workbook(BytesIO(body),read_only=True,data_only=True)
    except DataError: raise
    except Exception as exc: raise DataError('RBNZ_WORKBOOK_SCHEMA') from exc
    try:
        if not {'Data','Table Description','Series Definitions'}<=set(book.sheetnames):
            raise DataError('RBNZ_SHEETS_MISSING')
        for sheet in book:
            if (sheet.max_row or 0)>10_000 or (sheet.max_column or 0)>128:
                raise DataError('RBNZ_SHEET_BOUNDS')
            # The supplier currently declares A1:A1 for populated sheets.
            sheet.reset_dimensions()
        meta={r[0]:r[1] for r in book['Table Description'].iter_rows(max_row=40,max_col=2,values_only=True) if r[0]}
        published=meta.get('Published Date')
        if meta.get('Published By')!='Reserve Bank of New Zealand' or \
           meta.get('Table')!='Daily wholesale interest rates (% pa) - B2' or \
           not isinstance(published,(date,datetime)):
            raise DataError('RBNZ_PUBLICATION_SCHEMA')
        published=published.date() if isinstance(published,datetime) else published
        definitions=[r for r in book['Series Definitions'].iter_rows(max_row=128,max_col=5,values_only=True) if r[2]==series]
        group='Secondary market government bond closing yields'
        if len(definitions)!=1 or definitions[0][:4]!=(group,'5 year',series,'%pa'):
            raise DataError('RBNZ_DEFINITION_UNVERIFIED')
        rows=book['Data'].iter_rows(max_row=10_001,max_col=128,values_only=True)
        headers=[next(rows) for _ in range(5)]
        if headers[3][0]!='Unit' or headers[4][0]!='Series Id' or headers[4].count(series)!=1:
            raise DataError('RBNZ_SERIES_METADATA')
        col=headers[4].index(series)
        if (headers[0][col],headers[1][col],headers[3][col])!=(group,'5 year','%pa'):
            raise DataError('RBNZ_UNIT_OR_TENOR_UNVERIFIED')
        points=[];seen=set()
        for index,row in enumerate(rows,6):
            if index>10_000: raise DataError('RBNZ_SHEET_BOUNDS')
            day=row[0];value=row[col]
            if day is None and all(v is None for v in row): continue
            if not isinstance(day,(date,datetime)): raise DataError('RBNZ_DATE_SCHEMA')
            observed=(day.date() if isinstance(day,datetime) else day).isoformat()
            if date.fromisoformat(observed)>published: raise DataError('RBNZ_OBSERVATION_AFTER_PUBLICATION')
            if observed in seen: raise DataError('RBNZ_DUPLICATE_OBSERVATION')
            seen.add(observed)
            if value in (None,'','-','..','---'): continue
            if not finite(value): raise DataError('RBNZ_VALUE_SCHEMA')
            points.append((observed,float(value)))
        if not points: raise DataError('RBNZ_EMPTY')
        return points,published.isoformat()
    finally: book.close()

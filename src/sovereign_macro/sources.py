"""Source-specific parsing; series selection and maturity stay explicit."""
import csv
from datetime import timedelta, date, datetime
from io import StringIO, BytesIO
import re
import time
import zipfile
import calendar
from html.parser import HTMLParser
from urllib.parse import urljoin
from pypdf import PdfReader
import openpyxl
import xml.etree.ElementTree as ET
from .models import DataError, Observation, finite

def local(tag):
    return tag.rsplit('}',1)[-1]

def numeric(value):
    if value is None or value=='': return None
    try: value=float(value)
    except (TypeError,ValueError): return None
    return value if finite(value) else None

def parse_ecb(body):
    result={}
    for node in ET.fromstring(body).iter():
        if 'time' in node.attrib:
            day=node.attrib['time']
            date.fromisoformat(day)
            rates={child.attrib['currency']:float(child.attrib['rate']) for child in node if 'currency' in child.attrib}
            if any(not finite(v) or v<=0 for v in rates.values()): raise DataError('FX_RATE_CONTRACT')
            if day in result: raise DataError('DUPLICATE_FX_DATE')
            result[day]=rates
    if not result: raise DataError('FX_EMPTY')
    return result

def parse_canada(data,series):
    if series not in data.get('seriesDetail',{}): raise DataError('SERIES_METADATA_MISMATCH')
    label=data['seriesDetail'][series].get('label','').lower().replace('-',' ')
    if not re.search(r'\b5\s*year',label): raise DataError('CANADA_MATURITY_METADATA')
    return [(r['d'],v) for r in data['observations'] if (v:=numeric(r.get(series,{}).get('v'))) is not None]

def parse_riksbank(data):
    if not isinstance(data,list): raise DataError('RIKSBANK_SCHEMA')
    return [(r['date'][:10],v) for r in data if (v:=numeric(r.get('value'))) is not None]

def parse_treasury(body,tenor):
    rows=[]
    for node in ET.fromstring(body).iter():
        if local(node.tag)!='entry': continue
        fields={local(c.tag):c.text for c in node.iter()}
        value=numeric(fields.get(f'BC_{tenor}YEAR'))
        if value is not None and fields.get('NEW_DATE'): rows.append((fields['NEW_DATE'][:10],value))
    return rows

def parse_norway(body,tenor):
    reader=csv.DictReader(StringIO(body.decode('utf-8-sig')),delimiter=';')
    return [(r['TIME_PERIOD'],v) for r in reader if r['TENOR']==f'{tenor}Y' and r['INSTRUMENT_TYPE']=='GBON'
            and r['FREQ']=='B' and (v:=numeric(r['OBS_VALUE'])) is not None]

def parse_rba(body,series):
    rows=list(csv.reader(StringIO(body.decode('utf-8-sig'))))
    header=next((i for i,r in enumerate(rows) if series in r),None)
    if header is None: raise DataError('RBA_SERIES_MISSING')
    col=rows[header].index(series)
    units=next((r for r in rows[:header] if r and r[0]=='Units'),[])
    if len(units)<=col or units[col]!='Per cent per annum': raise DataError('RBA_UNIT_UNVERIFIED')
    result=[]
    for row in rows[header+1:]:
        if len(row)<=col: continue
        try: day=datetime.strptime(row[0],'%d-%b-%Y').date().isoformat()
        except ValueError: continue
        value=numeric(row[col])
        if value is not None: result.append((day,value))
    return result

def parse_japan(body,tenor):
    rows=list(csv.reader(StringIO(body.decode('cp932'))))
    if not rows or '%' not in ','.join(rows[0]): raise DataError('JAPAN_UNIT_UNVERIFIED')
    header=next((i for i,r in enumerate(rows) if f'{tenor}年' in r),None)
    if header is None: raise DataError('JAPAN_TENOR_MISSING')
    col=rows[header].index(f'{tenor}年')
    result=[]
    for row in rows[header+1:]:
        if len(row)<=col: continue
        match=re.fullmatch(r'([SHR])(\d+)\.(\d+)\.(\d+)',row[0].strip())
        if not match: continue
        era,y,m,d=match.groups()
        day=date({'S':1925,'H':1988,'R':2018}[era]+int(y),int(m),int(d))
        value=numeric(row[col])
        if value is not None: result.append((day.isoformat(),value))
    return result

def parse_kofia(body):
    rows=[]
    for node in ET.fromstring(body).iter():
        if local(node.tag)!='BISComDspDatDTO': continue
        fields={local(c.tag):c.text for c in node}
        day=fields.get('val1','') or ''
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',day): continue
        value=numeric(fields.get('val2'))
        if value is not None: rows.append((day,value))
    return rows

def parse_italy(body,series):
    rows=[]
    with zipfile.ZipFile(BytesIO(body)) as archive:
        for name in archive.namelist():
            if not name.endswith('_DATA.csv'): continue
            reader=csv.DictReader(StringIO(archive.read(name).decode('utf-8-sig')),delimiter=';')
            if series not in (reader.fieldnames or []): continue
            for row in reader:
                value=numeric(row[series])
                day=row.get('DATA_OSS','').replace('/','-')
                if value is not None: rows.append((day,value))
    return rows

def parse_belgium(body):
    rows=[]
    for node in ET.fromstring(body).iter():
        if local(node.tag)!='Obs': continue
        fields={local(c.tag):c.attrib for c in node}
        day=fields.get('ObsDimension',{}).get('value') or fields.get('Time',{}).get('value')
        val=fields.get('ObsValue',{}).get('value')
        # Generic SDMX uses Time text, compact SDMX uses attributes.
        if not day:
            day=next((c.text for c in node if local(c.tag)=='Time'),None)
        value=numeric(val)
        if day and value is not None: rows.append((day,value))
    return rows

def parse_spain(body,series):
    if series!='D_G0B1F0ZO': raise DataError('SPAIN_SERIES_CONTRACT')
    rows=list(csv.reader(StringIO(body.decode('cp1252'))))
    if not rows or rows[0].count(series)!=1: raise DataError('SPAIN_SERIES_MISSING')
    col=rows[0].index(series)
    metadata={r[0]:r for r in rows[:6] if r}
    def field(label):
        r=metadata.get(label,[])
        return r[col] if len(r)>col else ''
    description=field('DESCRIPCIÓN DE LA SERIE').lower()
    if 'bonos' not in description or 'del estado' not in description or not re.search(r'\b5 años\b',description):
        raise DataError('SPAIN_MATURITY_OR_ISSUER')
    if field('DESCRIPCIÓN DE LAS UNIDADES')!='Porcentaje' or field('FRECUENCIA')!='DIARIA':
        raise DataError('SPAIN_UNIT_OR_FREQUENCY')
    months={m:i+1 for i,m in enumerate(('ENE','FEB','MAR','ABR','MAY','JUN','JUL','AGO','SEP','OCT','NOV','DIC'))}
    result=[]
    for r in rows[6:]:
        if not r: continue
        if r[0] in ('FUENTE','NOTAS'): continue
        match=re.fullmatch(r'(\d{1,2}) ([A-Z]{3}) (\d{4})',r[0].strip())
        if not match: raise DataError('SPAIN_DATE_SCHEMA')
        d,m,y=match.groups()
        if m not in months or len(r)<=col: raise DataError('SPAIN_ROW_SCHEMA')
        day=date(int(y),months[m],int(d)).isoformat()
        if r[col] in ('','_'): continue
        value=numeric(r[col])
        if value is None: raise DataError('SPAIN_VALUE_SCHEMA')
        result.append((day,value))
    return result

def parse_slovakia(body,series):
    if series!='ZCY5Y': raise DataError('SLOVAKIA_SERIES_CONTRACT')
    w=openpyxl.load_workbook(BytesIO(body),read_only=True,data_only=True)
    try:
        if 'Yields_SK' not in w.sheetnames: raise DataError('SLOVAKIA_SHEET_SCHEMA')
        rows=iter(w['Yields_SK'].rows)
        header=next((r for r in rows if [c.value for c in r[:3]]==['YYYY','MM','DD']),None)
        if header is None: raise DataError('SLOVAKIA_DATE_SCHEMA')
        values=[c.value for c in header]
        if values.count(series)!=1: raise DataError('SLOVAKIA_MATURITY_SCHEMA')
        col=values.index(series); result=[]
        for r in rows:
            if all(c.value is None for c in r): continue
            if len(r)<=col: raise DataError('SLOVAKIA_ROW_SCHEMA')
            parts=[c.value for c in r[:3]]
            if not all(isinstance(v,(int,float)) and not isinstance(v,bool) and v==int(v) for v in parts):
                raise DataError('SLOVAKIA_DATE_SCHEMA')
            day=date(*(int(v) for v in parts)).isoformat()
            if r[col].value is None: continue
            if '%' in r[col].number_format: raise DataError('SLOVAKIA_UNIT_SCHEMA')
            value=numeric(r[col].value)
            if value is None: raise DataError('SLOVAKIA_VALUE_SCHEMA')
            result.append((day,value))
        return result
    finally: w.close()

def parse_iceland(body):
    w=openpyxl.load_workbook(BytesIO(body),read_only=True,data_only=True)
    try:
        if 'FLV' not in w.sheetnames: raise DataError('ICELAND_SHEET_SCHEMA')
        rows=list(w['FLV'].rows)
        header=next((i for i,r in enumerate(rows) if r and r[0].value=='Dagsetning (Date)'),None)
        if header is None or header==0: raise DataError('ICELAND_DATE_SCHEMA')
        groups=rows[header-1]; labels=[str(c.value or '') for c in groups]
        matches=[i for i,label in enumerate(labels) if '(Par-yield, nominal)' in label or label=='Par-yield, nominal']
        if len(matches)!=1: raise DataError('ICELAND_NOMINAL_PAR_SCHEMA')
        start=matches[0]
        end=next((i for i in range(start+1,len(groups)) if groups[i].value is not None),len(groups))
        cols=[i for i in range(start,end) if rows[header][i].value==5]
        if len(cols)!=1: raise DataError('ICELAND_MATURITY_SCHEMA')
        col=cols[0]
        correction=next((i for i,label in enumerate(labels) if i>0 and 'corrected calculation' in label),None)
        note=next((i for i,label in enumerate(labels) if label=='Athugasemd'),None)
        result=[]; notes={}
        for r in rows[header+1:]:
            observed=r[0].value
            if observed is None: continue
            if not isinstance(observed,(datetime,date)): raise DataError('ICELAND_DATE_SCHEMA')
            day=observed.date().isoformat() if isinstance(observed,datetime) else observed.isoformat()
            cell=r[col]
            if cell.value is None: continue
            # Only inspected Excel scaling formats; quoted/escaped % is literal.
            if cell.number_format not in ('0%','0.00%','#,##0.00%'):
                raise DataError('ICELAND_PERCENT_FORMAT_REQUIRED')
            value=numeric(cell.value)
            if value is None: raise DataError('ICELAND_VALUE_SCHEMA')
            result.append((day,value*100))
            flags='; '.join(f'{label}={r[i].value}' for label,i in [('correction',correction),('note_reference',note)]
                            if i is not None and r[i].value is not None)
            if day in notes and notes[day]!=flags: raise DataError('ICELAND_CONFLICTING_NOTES')
            notes[day]=flags
        return result,notes
    finally: w.close()

def parse_israel(body,series):
    if series!='ZC_TSB_ZND_05Y_MA': raise DataError('ISRAEL_SERIES_CONTRACT')
    matches=[n for n in ET.fromstring(body).iter() if local(n.tag)=='Series' and n.attrib.get('SERIES_CODE')==series]
    if len(matches)!=1: raise DataError('ISRAEL_SERIES_MISSING_OR_DUPLICATE')
    node=matches[0]
    contract={'FREQ':'M','NOMINAL_REAL':'N','DATA_TYPE':'ZC_YTM','TIME_TO_MATURITY':'Y05T05',
              'TIME_COLLECT':'A','DATA_SOURCE':'BOI_IS','UNIT_MULT':'0','UNIT_MEASURE':'PT'}
    if any(node.attrib.get(k)!=v for k,v in contract.items()): raise DataError('ISRAEL_DIMENSION_CONTRACT')
    rows=[]
    for obs in node:
        if local(obs.tag)!='Obs': continue
        period=obs.attrib.get('TIME_PERIOD','')
        if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period): raise DataError('ISRAEL_MONTH_SCHEMA')
        raw=obs.attrib.get('OBS_VALUE')
        if raw in (None,''): continue
        value=numeric(raw)
        if value is None: raise DataError('ISRAEL_VALUE_SCHEMA')
        rows.append((period,value))
    return rows

def select_latest_month(rows,as_of):
    # Only complete months; preserve the actual monthly period in Observation.
    eligible=[]
    for period,value in rows:
        if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period): raise DataError('MONTH_SCHEMA')
        y,m=map(int,period.split('-'))
        if date(y,m,calendar.monthrange(y,m)[1])<=as_of: eligible.append((period,value))
    parsed={}
    for period,value in eligible:
        if period in parsed and parsed[period]!=value: raise DataError('CONFLICTING_DUPLICATE_YIELD')
        parsed[period]=value
    if not parsed: raise DataError('YIELD_EMPTY_AS_OF')
    period=max(parsed)
    return period,parsed[period]

def select_latest(rows,as_of):
    parsed={}
    for period,value in rows:
        observed=date.fromisoformat(period)
        if observed>as_of: continue
        if period in parsed and parsed[period]!=value: raise DataError('CONFLICTING_DUPLICATE_YIELD')
        parsed[period]=value
    if not parsed: raise DataError('YIELD_EMPTY_AS_OF')
    period=max(parsed)
    return period,parsed[period]

CNB_BULLETIN_INDEX='https://www.cnb.cz/en/statistics/money_and_banking_stat/monetary-statistics-monthly-bulletin/index.html'

def discover_czech_bulletin(body,as_of):
    class Links(HTMLParser):
        def __init__(self):
            super().__init__(); self.links=[]
        def handle_starttag(self,tag,attrs):
            if tag=='a': self.links.append(dict(attrs).get('href',''))
    parser=Links(); parser.feed(body.decode('utf-8'))
    candidates=set()
    for link in parser.links:
        url=urljoin(CNB_BULLETIN_INDEX,link)
        match=re.fullmatch(r'https://www\.cnb\.cz/export/sites/cnb/en/statistics/\.galleries/money_and_banking_stat/mon_bank_stat/(\d{4})/menstat_(\d{4})-(\d{2})_EN\.pdf',url)
        if not match: continue
        folder,y,m=match.groups()
        if folder!=y or not 1<=int(m)<=12: continue
        edition=f'{y}-{m}'
        if edition<=as_of.strftime('%Y-%m'): candidates.add((edition,url))
    if not candidates: raise DataError('CZE_PUBLISHED_EDITION_MISSING')
    edition,url=max(candidates)
    return url,edition

def parse_czech_pages(pages,edition):
    # The PDF table is drawn with positioned glyphs. Its narrative explicitly
    # labels the current month and 5Y yield; never infer a value from glyph order.
    pages=[re.sub(r'\s+',' ',p.replace('\u2212','-')).strip() for p in pages]
    table=[p for p in pages if '1.3 TABLE 2B – CAPITAL MARKET INTEREST RATES' in p]
    commentary=[p for p in pages if '1.4 COMMENTARY ON TABLES 1 – 2' in p]
    if len(table)!=1 or len(commentary)!=1: raise DataError('CZE_SECTION_SCHEMA')
    section=table[0].split('1.3 TABLE 2B – CAPITAL MARKET INTEREST RATES',1)[1]
    if not re.search(r'\b5 years\b',section) or not all(s in section for s in ('(in %, monthly average)','Bond yields','Source: Czech National Bank.')):
        raise DataError('CZE_TABLE_METADATA')
    for page in (table[0],commentary[0]):
        editions=re.findall(r'Monetary Statistics – (\d{1,2})/(\d{4})',page)
        if len(editions)!=1 or f'{editions[0][1]}-{int(editions[0][0]):02}'!=edition:
            raise DataError('CZE_EDITION_MISMATCH')
    text=commentary[0]
    matches=re.findall(r'Commentary on key interest rates \(Table 1\) and financial market interest rates \(Table 2\): ([A-Za-z]+) (\d{4})\.',text)
    if len(matches)!=1: raise DataError('CZE_REFERENCE_MONTH_SCHEMA')
    month,y=matches[0]
    months={calendar.month_name[m]:m for m in range(1,13)}
    if month not in months: raise DataError('CZE_REFERENCE_MONTH_SCHEMA')
    period=f'{y}-{months[month]:02}'
    lag=(int(edition[:4])-int(y))*12+int(edition[5:])-months[month]
    if not 0<=lag<=3: raise DataError('CZE_REFERENCE_MONTH_LAG')
    text=text.split('1.4.2 FINANCIAL MARKET INTEREST RATES',1)
    if len(text)!=2: raise DataError('CZE_COMMENTARY_SCHEMA')
    # Split at sentence-ending periods, not decimal points. Missing 5Y data
    # must never consume a later 10Y sentence or its percentage.
    sentences=[s.strip() for s in re.split(r'\.(?:\s+|$)',text[1]) if re.search(r'\byield on the 5Y bond\b',s)]
    if len(sentences)!=1: raise DataError('CZE_FIVE_YEAR_AMBIGUOUS')
    number=r'-?\d+(?:\.\d+)?'
    value=re.fullmatch(r'(?:The )?yield on the 5Y bond (?:increased|decreased|rose|fell)(?: by '+number+r' percentage points?)? to ('+number+r')%',sentences[0])
    if not value: raise DataError('CZE_YIELD_SCHEMA')
    return [(period,float(value.group(1)))]

def collect_yield(client,country,as_of,tenor=5):
    route=country['yield']; adapter=route['adapter']; series=route['series']
    if tenor not in (5,10): raise DataError('TENOR_NOT_IMPLEMENTED')
    if tenor!=5 and adapter not in ('norway','treasury','japan'):
        raise DataError('TENOR_NOT_IMPLEMENTED')
    start=(as_of-timedelta(days=21)).isoformat()
    body=None; notes={}; frequency='daily'
    if adapter=='canada':
        if tenor!=5: raise DataError('TENOR_NOT_IMPLEMENTED')
        url=f'https://www.bankofcanada.ca/valet/observations/{series}/json?start_date={start}&end_date={as_of}'
        parser=lambda p:parse_canada(p.json(),series)
    elif adapter=='norway':
        url=f'https://data.norges-bank.no/api/data/GOVT_GENERIC_RATES/?startPeriod={start}&format=csv'
        parser=lambda p:parse_norway(p.body,tenor)
    elif adapter=='riksbank':
        if not hasattr(client,'riksbank_series'):
            meta=client.fetch('https://api.riksbank.se/swea/v1/Series').json()
            client.riksbank_series={r['seriesId']:r for r in meta}
        metadata=client.riksbank_series.get(series,{})
        if 'maturity 5 years' not in metadata.get('longDescription','') or metadata.get('seriesClosed') is not False:
            raise DataError('RIKSBANK_SERIES_METADATA')
        if tenor!=5: raise DataError('TENOR_NOT_IMPLEMENTED')
        elapsed=time.monotonic()-getattr(client,'riksbank_last_request',0)
        if elapsed<1: time.sleep(1-elapsed)
        client.riksbank_last_request=time.monotonic()
        url=f'https://api.riksbank.se/swea/v1/Observations/{series}/{start}/{as_of}'
        parser=lambda p:parse_riksbank(p.json())
    elif adapter=='rba':
        url='https://www.rba.gov.au/statistics/tables/csv/f2-data.csv'
        parser=lambda p:parse_rba(p.body,series)
    elif adapter=='treasury':
        url=f'https://home.treasury.gov/resource-center/data-chart-center/interest-rates/pages/xml?data=daily_treasury_yield_curve&field_tdr_date_value={as_of.year}'
        series=f'BC_{tenor}YEAR'
        parser=lambda p:parse_treasury(p.body,tenor)
    elif adapter=='japan':
        url='https://www.mof.go.jp/jgbs/reference/interest_rate/jgbcm.csv'
        series=f'{tenor}年'
        parser=lambda p:parse_japan(p.body,tenor)
    elif adapter=='kofia':
        url='https://www.kofiabond.or.kr/proframeWeb/XMLSERVICES/'
        body=f'<message><proframeHeader><pfmAppName>BIS-KOFIABOND</pfmAppName><pfmSvcName>BISLastAskPrcROPSrchSO</pfmSvcName><pfmFnName>listTrm</pfmFnName></proframeHeader><systemHeader/><BISComDspDatDTO><val1>DD</val1><val2>{start.replace("-","")}</val2><val3>{as_of.isoformat().replace("-","")}</val3><val4>1530</val4><val5>{series}</val5></BISComDspDatDTO></message>'
        parser=lambda p:parse_kofia(p.body)
    elif adapter=='italy':
        url='https://a2a.bancaditalia.it/infostat/dataservices/export/EN/CSV/ALL/CUBE/BANKITALIA/DIFF/BMK0200'
        parser=lambda p:parse_italy(p.body,series)
    elif adapter=='belgium':
        url=f'https://nsidisseminate-stat.nbb.be/rest/data/BE2,DF_IROLOBE2,1.0/D.5Y.F?startPeriod={start}'
        parser=lambda p:parse_belgium(p.body)
    elif adapter=='spain':
        url='https://www.bde.es/webbe/es/estadisticas/compartido/datos/csv/ti_1_3.csv'
        parser=lambda p:parse_spain(p.body,series)
    elif adapter=='slovakia':
        url='https://nbs.sk/dokument/b912f986-f5ab-4a02-9033-97976d6dc849/stiahnut?force=false'
        parser=lambda p:parse_slovakia(p.body,series)
    elif adapter=='iceland':
        if series!='FLV:nominal:par:5Y': raise DataError('ICELAND_SERIES_CONTRACT')
        url='https://sedlabanki.is/library?itemid=4b7a7e67-a647-4e98-9190-0c1ca772179f'
        def parser(p):
            rows,annotations=parse_iceland(p.body)
            notes.update(annotations)
            return rows
    elif adapter=='czech':
        if series!='TABLE_2B:5Y:monthly_average': raise DataError('CZE_SERIES_CONTRACT')
        index=client.fetch(CNB_BULLETIN_INDEX)
        url,edition=discover_czech_bulletin(index.body,as_of)
        parser=lambda p:parse_czech_pages([page.extract_text() or '' for page in PdfReader(BytesIO(p.body)).pages],edition)
        frequency='monthly'
        notes['edition']=edition
    elif adapter=='israel':
        # Three calendar months, so monthly series are not filtered out by a daily window.
        month_index=as_of.year*12+as_of.month-1-3
        monthly_start=date(month_index//12,month_index%12+1,1).isoformat()
        url=f'https://edge.boi.gov.il/FusionEdgeServer/sdmx/v2/data/dataflow/BOI.STATISTICS/ZCM/1.0/{series}?startPeriod={monthly_start}'
        parser=lambda p:parse_israel(p.body,series)
        frequency='monthly'
    else: raise DataError(route.get('reason','ADAPTER_UNAVAILABLE'))
    p=client.fetch(url,method='POST' if body else 'GET',body=body)
    period,value=(select_latest_month if frequency=='monthly' else select_latest)(parser(p),as_of)
    return Observation(country['iso3'],f'yield_{tenor}y',value,period,country['currency'],
                       unit='unverified' if adapter=='belgium' else 'percent',
                       provider=adapter,dataset=adapter,series=series,source_date=p.source_date,
                       retrieved_at=p.retrieved_at,url=p.url,raw_sha256=p.sha256,
                       tenor_years=tenor,yield_type=route['yield_type'],redistribution=route['redistribution'],
                       frequency=frequency,notes=('publication_edition='+notes['edition']) if adapter=='czech' else notes.get(period,''))

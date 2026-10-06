"""UK and Swedish transaction diagnostics with source-specific definitions."""
from calendar import monthrange
from datetime import date, datetime
from io import BytesIO
import re
from urllib.parse import urljoin, urlparse
from zipfile import ZipFile

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from .models import DataError
from .liquidity import diagnostic, number

DMO_PAGE='https://www.dmo.gov.uk/data/gilt-market/turnover-data/'
RIKSBANK_PAGE='https://www.riksbank.se/en-gb/statistics/turnover-statistics/fi-market/'
RIKSBANK_DESCRIPTION='https://www.riksbank.se/globalassets/media/statistik/omsattningsstatistik/penning_obligationsmarknad/daily_fi_description.xlsx'


def workbook(body, prefix, limit):
    try:
        with ZipFile(BytesIO(body)) as archive:
            if len(archive.infolist())>1000 or sum(i.file_size for i in archive.infolist())>limit:
                raise DataError(prefix+'_WORKBOOK_TOO_LARGE')
        return load_workbook(BytesIO(body),read_only=True,data_only=True)
    except DataError: raise
    except Exception as exc: raise DataError(prefix+'_WORKBOOK_SCHEMA') from exc


def discover_workbook(body, page, label):
    if len(body)>2_000_000: raise DataError('LIQUIDITY_INDEX_TOO_LARGE')
    soup=BeautifulSoup(body,'html.parser')
    links=set()
    for anchor in soup.select('a[href]'):
        text=' '.join(anchor.get_text(' ',strip=True).replace('’',"'").split())
        if label not in text: continue
        url=urljoin(page,anchor['href']);parsed=urlparse(url)
        if parsed.scheme!='https' or parsed.netloc!=urlparse(page).netloc or \
           not parsed.path.endswith('.xlsx') or parsed.query or parsed.fragment:
            raise DataError('LIQUIDITY_WORKBOOK_LINK_UNVERIFIED')
        links.add(url)
    if len(links)!=1: raise DataError('LIQUIDITY_WORKBOOK_LINK_UNVERIFIED')
    return links.pop()


def parse_dmo_liquidity(body, as_of, max_age_days=185):
    book=workbook(body,'DMO',10_000_000)
    try:
        if 'Notes' not in book.sheetnames: raise DataError('DMO_DEFINITION_UNVERIFIED')
        notes=' '.join(str(v or '') for row in book['Notes'].iter_rows(max_row=30,max_col=2,values_only=True) for v in row)
        for phrase in ('market values','sum of all trade reports','neither counterparty is a GEMM',
                       'total purchases and sales','do not include turnover on repo'):
            if phrase not in notes: raise DataError('DMO_DEFINITION_UNVERIFIED')
        points={}
        years=sorted((n for n in book.sheetnames if re.fullmatch(r'\d{4}',n) and int(n)<=as_of.year),reverse=True)
        for name in years:
            sheet=book[name]
            if sheet.max_row is None or sheet.max_column is None or sheet.max_row>10_000 or sheet.max_column>32:
                raise DataError('DMO_SHEET_BOUNDS')
            rows=list(sheet.iter_rows(max_row=100,max_col=7,values_only=True))
            if rows[0][0]!='ALL GEMMs turnover by maturity band' or rows[1][0]!=int(name):
                raise DataError('DMO_SHEET_SCHEMA')
            observed=None;period=None;parts=[];bands=set();header=False;closed=True
            for row in rows[2:]:
                label=str(row[0] or '').strip()
                match=re.fullmatch(r'Q([1-4]) \([^)]*\)',label)
                if match:
                    if not closed: raise DataError('DMO_TOTAL_MISSING')
                    q=int(match[1]);month=q*3
                    observed=date(int(name),month,monthrange(int(name),month)[1]);period=f'{name}-Q{q}'
                    parts=[];bands=set();header=False;closed=False
                elif observed is not None and label=='BAND':
                    expected=('Customer','Customer','Professional','Professional','Combined','Combined')
                    if any(not str(v or '').startswith(e) for v,e in zip(row[1:],expected)) or \
                       any('£ million' not in str(row[i]) for i in (1,3,5)):
                        raise DataError('DMO_UNIT_UNVERIFIED')
                    header=True
                elif observed is not None and label:
                    if label.startswith('Note:'): continue
                    if not header: raise DataError('DMO_SHEET_SCHEMA')
                    if closed: raise DataError('DMO_DUPLICATE_TOTAL')
                    values=[number(row[i]) for i in (1,3,5)]
                    if abs(values[0]+values[1]-values[2])>0.01: raise DataError('DMO_TOTAL_MISMATCH')
                    if label=='Total':
                        if not parts or any(abs(sum(p[i] for p in parts)-values[i])>0.01 for i in range(3)):
                            raise DataError('DMO_TOTAL_MISMATCH')
                        if observed<=as_of:
                            if observed in points: raise DataError('DMO_DUPLICATE_OBSERVATION')
                            points[observed]=(period,values[2])
                        closed=True
                    else:
                        if not re.fullmatch(r'(?:\d+[-–]\d+ (?:years|yrs|months|mths)|\d+\+ (?:years|yrs)|[Ii]ndex-linked|Strips)',label):
                            raise DataError('DMO_BAND_SCHEMA')
                        if label.lower() in bands: raise DataError('DMO_DUPLICATE_BAND')
                        bands.add(label.lower())
                        parts.append(values)
            if not closed: raise DataError('DMO_TOTAL_MISSING')
            if points: break
        return diagnostic(points,as_of,max_age_days,provider='uk_dmo',frequency='quarterly',
            unit='GBP_million_quarterly_purchase_plus_sale',
            definition='GEMM-reported gilt purchases plus sales at market value; all maturities including index-linked; '
                       'repo and derivatives excluded; non-GEMM/non-GEMM trades absent; reporter sums can double-count; '
                       'not unique-trade volume or a single-5Y measure')
    finally: book.close()


def verify_riksbank_description(body):
    book=workbook(body,'RIKSBANK',1_000_000)
    try:
        if 'Description' not in book.sheetnames: raise DataError('RIKSBANK_DEFINITION_UNVERIFIED')
        rows=list(book['Description'].iter_rows(max_row=60,max_col=4,values_only=True))
        mapping={r[1]:str(r[2] or '') for r in rows if r[1] is not None}
        if mapping.get('GVB')!='Government Bonds' or mapping.get('SP')!='Spot' or mapping.get('RE')!='Repo' or \
           'Secondary Market is the sum of all other counterparts' not in mapping.get('PRIMM','') or \
           'Adjusted for double-counting' not in mapping.get('Adjusted amount','') or \
           'dividing those amounts by two' not in mapping.get('Adjusted amount','') or \
           not any('Millions SEK' in str(v) for r in rows for v in r):
            raise DataError('RIKSBANK_DEFINITION_UNVERIFIED')
    finally: book.close()


def parse_riksbank_liquidity(body, as_of, max_age_days=7):
    book=workbook(body,'RIKSBANK',80_000_000)
    try:
        if 'Data' not in book.sheetnames: raise DataError('RIKSBANK_DATA_SHEET_MISSING')
        sheet=book['Data']
        if sheet.max_row is None or sheet.max_column is None or sheet.max_row>300_000 or sheet.max_column!=7:
            raise DataError('RIKSBANK_SHEET_BOUNDS')
        rows=sheet.iter_rows(max_col=7,values_only=True)
        if next(rows)!=('Trade date','Asset','Contract','Counterparty','Amount (million SEK)',
                        'Adjusted amount (million SEK)','Complete'):
            raise DataError('RIKSBANK_UNIT_UNVERIFIED')
        points={};keys=set();complete={}
        for day,asset,contract,counterparty,amount,adjusted,flag in rows:
            if asset!='GVB' or contract!='SP' or counterparty=='PRIMM': continue
            if not isinstance(day,(date,datetime)): raise DataError('RIKSBANK_DATE_SCHEMA')
            observed=day.date() if isinstance(day,datetime) else day
            if observed>as_of: continue
            if counterparty not in {'REP','OMM','BROKNO','CUSE','CUFO','RIX'}:
                raise DataError('RIKSBANK_COUNTERPARTY_SCHEMA')
            key=(observed,counterparty)
            if key in keys: raise DataError('RIKSBANK_DUPLICATE_OBSERVATION')
            keys.add(key)
            value=number(adjusted);raw=number(amount)
            if value>raw: raise DataError('RIKSBANK_ADJUSTED_AMOUNT_SCHEMA')
            if flag is not True and flag is not False and flag is not None:
                raise DataError('RIKSBANK_COMPLETENESS_SCHEMA')
            complete[observed]=complete.get(observed,True) and flag is True
            points[observed]=(observed.isoformat(),points.get(observed,('',0))[1]+value)
        if points and not complete[max(points)]: raise DataError('RIKSBANK_INCOMPLETE_LATEST_OBSERVATION')
        return diagnostic(points,as_of,max_age_days,provider='riksbank',frequency='daily',
            unit='SEK_million_daily_adjusted_turnover',
            definition='SELMA reporting-counterparty sample; GVB government bonds, SP spot; excludes primary market '
                       '(PRIMM), repo, forwards, inflation-linked bonds and bills; adjusted amount already corrects '
                       'reporter double-counting; all maturities, not a single-5Y or whole-market measure')
    finally: book.close()


def collect_europe_liquidity(client, settings, as_of):
    result={}
    sources=[('GBR','uk_dmo',DMO_PAGE,"All GEMMs' turnover by maturity band",parse_dmo_liquidity,
              settings.get('quarterly_max_age_days',185),'uk_dmo_enabled'),
             ('SWE','riksbank',RIKSBANK_PAGE,'Fixed Income, Daily Turnover',parse_riksbank_liquidity,
              settings.get('daily_max_age_days',7),'riksbank_enabled')]
    for iso,provider,page,label,parser,age,flag in sources:
        if not settings.get(flag,True): continue
        try:
            index=client.fetch(page)
            if index.url!=page: raise DataError('LIQUIDITY_UNEXPECTED_REDIRECT')
            url=discover_workbook(index.body,page,label)
            definition=None
            if iso=='SWE':
                definition=client.fetch(RIKSBANK_DESCRIPTION)
                if definition.url!=RIKSBANK_DESCRIPTION: raise DataError('LIQUIDITY_UNEXPECTED_REDIRECT')
                verify_riksbank_description(definition.body)
            payload=client.fetch(url)
            if payload.url!=url: raise DataError('LIQUIDITY_UNEXPECTED_REDIRECT')
            row=parser(payload.body,as_of,age)
            row.update(url=url,raw_sha256=payload.sha256,retrieved_at=payload.retrieved_at,
                       transport=payload.transport,index_url=page,index_sha256=index.sha256,
                       redistribution=settings.get('redistribution','pending'))
            if definition: row.update(definition_url=definition.url,definition_sha256=definition.sha256)
            result[iso]=row
        except Exception as exc:
            result[iso]=dict(value=None,status='UNAVAILABLE',usable_for_scoring=False,provider=provider,
                            url=page,reason=type(exc).__name__+':'+str(exc)[:180])
    return result

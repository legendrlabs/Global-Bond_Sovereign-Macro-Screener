"""Source-specific transaction diagnostics, never a comparable liquidity score."""
from calendar import monthrange
from datetime import date
from io import BytesIO
import json
import re
from zipfile import ZipFile

from openpyxl import load_workbook
from .models import DataError, finite

NYFED_SERIES='PDGSWOEXTTOT'
NYFED_BREAK='SBN2024'
NYFED_CATALOG='https://markets.newyorkfed.org/api/pd/list/timeseries.json'
NYFED_LATEST='https://markets.newyorkfed.org/api/pd/latest/SBN2024.json'
JSDA_WORKBOOK='https://www.jsda.or.jp/shiryoshitsu/toukei/tentoubaibai/koushasai.xlsx'


def number(value):
    if isinstance(value,bool): raise DataError('LIQUIDITY_VALUE_SCHEMA')
    try: result=float(value)
    except (ValueError,TypeError) as exc: raise DataError('LIQUIDITY_VALUE_SCHEMA') from exc
    if not finite(result) or result<0: raise DataError('LIQUIDITY_VALUE_SCHEMA')
    return result


def diagnostic(points,as_of,max_age_days,**metadata):
    if not points: raise DataError('LIQUIDITY_NO_COMPLETED_OBSERVATION')
    observed=max(points);period,value=points[observed];age=(as_of-observed).days
    return dict(metadata,value=value if age<=max_age_days else None,period=period,
        observation_date=observed.isoformat(),age_days=age,
        status='AVAILABLE_DIAGNOSTIC' if age<=max_age_days else 'STALE',
        usable_for_scoring=False,reason='Source-specific transaction volume; not harmonized across countries')


def parse_nyfed_liquidity(body,as_of,max_age_days=30):
    try: rows=json.loads(body)['pd']['timeseries']
    except (ValueError,KeyError,TypeError) as exc: raise DataError('NYFED_JSON_SCHEMA') from exc
    if not isinstance(rows,list): raise DataError('NYFED_JSON_SCHEMA')
    points={}
    for row in rows:
        if not isinstance(row,dict): raise DataError('NYFED_JSON_SCHEMA')
        if row.get('keyid')!=NYFED_SERIES: continue
        try: observed=date.fromisoformat(row['asofdate'])
        except (ValueError,TypeError,KeyError) as exc: raise DataError('NYFED_DATE_SCHEMA') from exc
        if observed>as_of: continue
        if observed in points: raise DataError('NYFED_DUPLICATE_OBSERVATION')
        points[observed]=(observed.isoformat(),number(row.get('value')))
    return diagnostic(points,as_of,max_age_days,provider='nyfed',series=NYFED_SERIES,
        frequency='weekly_daily_average',unit='USD_million_daily_average',
        definition='Primary-dealer Treasury outright purchases plus sales excluding TIPS; all maturities; '
                   'inter-dealer brokers plus others; trading-day average; includes when-issued/forward delivery; '
                   'repo excluded; dealer reporting can double-count interdealer activity')


def verify_nyfed_catalog(body):
    try: rows=json.loads(body)['pd']['timeseries']
    except (ValueError,KeyError,TypeError) as exc: raise DataError('NYFED_CATALOG_SCHEMA') from exc
    if not isinstance(rows,list) or any(not isinstance(r,dict) for r in rows):
        raise DataError('NYFED_CATALOG_SCHEMA')
    selected=[r for r in rows if r.get('seriesbreak')==NYFED_BREAK and r.get('keyid')==NYFED_SERIES]
    if len(selected)!=1: raise DataError('NYFED_DEFINITION_UNVERIFIED')
    description=' '.join(str(selected[0].get('description','')).upper().split())
    required=('TREASURY SECURITIES (EXCLUDING TREASURY INFLATION-PROTECTED SECURITIES (TIPS))',
              'DEALER TRANSACTIONS WITH INTER-DEALER BROKERS','DEALER TRANSACTIONS WITH OTHERS')
    if any(t not in description for t in required) or 'CHANGE' in description or 'POSITION' in description:
        raise DataError('NYFED_DEFINITION_UNVERIFIED')


def parse_jsda_liquidity(body,as_of,max_age_days=120):
    try:
        with ZipFile(BytesIO(body)) as archive:
            if len(archive.infolist())>1000 or sum(f.file_size for f in archive.infolist())>30_000_000:
                raise DataError('JSDA_WORKBOOK_TOO_LARGE')
        book=load_workbook(BytesIO(body),read_only=True,data_only=True)
    except DataError: raise
    except Exception as exc: raise DataError('JSDA_WORKBOOK_SCHEMA') from exc
    try:
        if '(Ｂ)一般売買高' not in book.sheetnames: raise DataError('JSDA_OUTRIGHT_SHEET_MISSING')
        sheet=book['(Ｂ)一般売買高']
        if sheet.max_row is None or sheet.max_column is None or sheet.max_row>5000 or sheet.max_column>128:
            raise DataError('JSDA_SHEET_BOUNDS')
        unit=str(sheet.cell(1,1).value or '')
        if '100 million yen' not in unit: raise DataError('JSDA_UNIT_UNVERIFIED')
        if 'Sell＋Purchase' not in str(sheet.cell(1,4).value or '') or \
           'Outright transactions' not in str(sheet.cell(2,4).value or '') or \
           'Government Bonds' not in str(sheet.cell(3,4).value or ''):
            raise DataError('JSDA_DEFINITION_UNVERIFIED')
        points={}
        for row in sheet.iter_rows(min_row=5,max_col=4,values_only=True):
            if len(row)<4 or row[2]!='Total' or row[1]!='合計': continue
            label=str(row[0]).strip()
            if re.fullmatch(r'\d{4}(?:年度|計)',label): continue
            match=re.fullmatch(r'(\d{4})/(\d{2})',label)
            if not match: raise DataError('JSDA_MONTH_SCHEMA')
            try:
                year,month=map(int,match.groups());observed=date(year,month,monthrange(year,month)[1])
            except ValueError as exc: raise DataError('JSDA_MONTH_SCHEMA') from exc
            if observed>as_of: continue
            if observed in points: raise DataError('JSDA_DUPLICATE_OBSERVATION')
            points[observed]=(f'{year:04}-{month:02}',number(row[3]))
        return diagnostic(points,as_of,max_age_days,provider='jsda',frequency='monthly',
            unit='JPY_100_million_monthly_sell_plus_purchase',
            definition='JSDA regular plus special members; government-bond OTC outright sell plus purchase; '
                       'face amount; includes Treasury bills, when-issued/GX bonds, auctions and BOJ activity; '
                       'excludes exchange trades and repo; not a unique-trade or single-5Y measure')
    finally: book.close()


def collect_liquidity(client,settings,as_of):
    result={}
    for iso,url,parser,age in [('USA',NYFED_LATEST,parse_nyfed_liquidity,settings.get('weekly_max_age_days',30)),
                              ('JPN',JSDA_WORKBOOK,parse_jsda_liquidity,settings.get('monthly_max_age_days',120))]:
        try:
            catalog=None
            if iso=='USA':
                catalog=client.fetch(NYFED_CATALOG)
                if catalog.url!=NYFED_CATALOG: raise DataError('NYFED_UNEXPECTED_REDIRECT')
                verify_nyfed_catalog(catalog.body)
            payload=client.fetch(url)
            if payload.url!=url: raise DataError('LIQUIDITY_UNEXPECTED_REDIRECT')
            row=parser(payload.body,as_of,age)
            row.update(url=payload.url,raw_sha256=payload.sha256,retrieved_at=payload.retrieved_at,
                transport=payload.transport,redistribution=settings.get('redistribution','pending'))
            if catalog: row.update(definition_url=catalog.url,definition_sha256=catalog.sha256)
            result[iso]=row
        except Exception as exc:
            result[iso]=dict(value=None,status='UNAVAILABLE',usable_for_scoring=False,
                provider='nyfed' if iso=='USA' else 'jsda',url=url,
                reason=type(exc).__name__+':'+str(exc)[:180])
    return result
